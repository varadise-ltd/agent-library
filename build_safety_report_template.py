from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"C:\Users\zecalee\Documents\ChatGPT\Varadise Agent Library\Bilingual_Safety_Report_Template_EN_TC.docx"

NAVY = "183B56"; TEAL = "1C7C7D"; PALE = "EAF4F4"; LIGHT = "F3F6F8"
MID = "D7E1E7"; TEXT = "20313C"; MUTED = "5D6B75"; WHITE = "FFFFFF"
AMBER = "FFF4D6"; RED = "FCE8E6"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+tag))
        if node is None: node=OxmlElement('w:'+tag); tcMar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')

def set_cell_width(cell, dxa):
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
    if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'),str(dxa)); tcW.set(qn('w:type'),'dxa')

def table_geometry(table, widths, indent=120):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    tblPr=table._tbl.tblPr
    tblW=tblPr.find(qn('w:tblW'))
    if tblW is None: tblW=OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'),str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    tblInd=tblPr.find(qn('w:tblInd'))
    if tblInd is None: tblInd=OxmlElement('w:tblInd'); tblPr.append(tblInd)
    tblInd.set(qn('w:w'),str(indent)); tblInd.set(qn('w:type'),'dxa')
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(w)); grid.append(gc)
    for row in table.rows:
        for i,c in enumerate(row.cells): set_cell_width(c,widths[i]); margins(c)

def borders(table, color=MID, size='5'):
    tblPr=table._tbl.tblPr; b=tblPr.find(qn('w:tblBorders'))
    if b is None: b=OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),size); el.set(qn('w:color'),color); b.append(el)

def keep_row(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)

def mark_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def font(run, size=9.5, bold=False, color=TEXT, italic=False):
    run.font.name='Arial'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=RGBColor.from_string(color)

def cell_text(cell, text, bold=False, color=TEXT, size=9.2, align=None):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
    if align is not None: p.alignment=align
    font(p.add_run(text),size,bold,color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_field_table(doc, rows, widths=(2750,6610)):
    t=doc.add_table(rows=0, cols=2); table_geometry(t,list(widths)); borders(t)
    for label,value in rows:
        cells=t.add_row().cells; keep_row(t.rows[-1]); shade(cells[0],LIGHT)
        cell_text(cells[0],label,True,NAVY); cell_text(cells[1],value)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def add_matrix(doc, headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers)); table_geometry(t,widths); borders(t)
    for i,h in enumerate(headers): shade(t.rows[0].cells[i],NAVY); cell_text(t.rows[0].cells[i],h,True,WHITE,8.6,WD_ALIGN_PARAGRAPH.CENTER)
    keep_row(t.rows[0]); mark_header(t.rows[0])
    for row in rows:
        cells=t.add_row().cells; keep_row(t.rows[-1])
        for i,v in enumerate(row):
            if len(t.rows)%2==1: shade(cells[i],'FAFBFC')
            cell_text(cells[i],v,size=8.7)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def heading(doc, en, zh):
    p=doc.add_paragraph(style='Heading 1'); p.paragraph_format.keep_with_next=True
    r=p.add_run(en+'  |  '+zh); font(r,15,True,NAVY)

def subheading(doc,en,zh):
    p=doc.add_paragraph(style='Heading 2'); p.paragraph_format.keep_with_next=True
    font(p.add_run(en+'  |  '+zh),11.5,True,TEAL)

def note(doc, text, fill=PALE):
    t=doc.add_table(rows=1,cols=1); table_geometry(t,[9360]); borders(t,TEAL,'7'); shade(t.cell(0,0),fill); cell_text(t.cell(0,0),text,size=9)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def page_break(doc): doc.add_page_break()

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(1); sec.right_margin=Inches(1); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)

styles=doc.styles
normal=styles['Normal']; normal.font.name='Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); normal.font.size=Pt(9.5); normal.font.color.rgb=RGBColor.from_string(TEXT)
normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.12
for name,size,before,after,color in [('Heading 1',15,14,7,NAVY),('Heading 2',11.5,9,4,TEAL),('Heading 3',10.5,7,3,NAVY)]:
    s=styles[name]; s.font.name='Arial'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

# Header/footer
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(hp.add_run('SAFETY REPORT / 安全報告'),8.5,True,MUTED)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run('Controlled template / 受控範本     |     Page / 頁碼 '),8,color=MUTED)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(16); p.paragraph_format.space_after=Pt(3); font(p.add_run('WORKPLACE & SITE'),10,True,TEAL)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); font(p.add_run('SAFETY REPORT TEMPLATE'),25,True,NAVY)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(18); font(p.add_run('工作場所及工地安全報告範本'),18,True,TEAL)
note(doc,'Use this form to report an incident, injury, illness, near miss, unsafe condition, property damage, or environmental event. Complete promptly and attach supporting evidence.\n本表格適用於事故、受傷、疾病、險失事件、不安全情況、財物損壞或環境事件。請盡快填寫，並附上相關證據。')
add_field_table(doc,[
('Organisation / 機構','____________________________________________________________'),
('Project / Site / 項目／地點','____________________________________________________________'),
('Report reference / 報告編號','__________________________'),
('Report status / 報告狀態','☐ Draft 草稿    ☐ Final 定稿    ☐ Closed 已結案'),
('Confidentiality / 保密級別','☐ Internal 內部    ☐ Restricted 受限    ☐ Other 其他: __________'),
('Prepared by / 編製人','__________________________    Date / 日期: ____ / ____ / ______'),
('Reviewed by / 覆核人','__________________________    Date / 日期: ____ / ____ / ______'),
])
subheading(doc,'Report type','報告類別')
add_matrix(doc,['Select / 選擇','Category / 類別','Select / 選擇','Category / 類別'],[
('☐','Injury / illness\n受傷／疾病','☐','Near miss\n險失事件'),('☐','Unsafe act / condition\n不安全行為／情況','☐','Property damage\n財物損壞'),('☐','Environmental event\n環境事件','☐','Emergency / evacuation\n緊急事故／疏散'),('☐','Other / 其他: __________','☐','Regulatory notification\n法定呈報')],[600,4080,600,4080])

page_break(doc); heading(doc,'1. EVENT DETAILS','1. 事件資料')
add_field_table(doc,[
('Date of event / 事件日期','____ / ____ / ______'),('Time / 時間','________  ☐ 24-hour / 24小時制'),('Exact location / 確實地點','____________________________________________________________'),('Reported date & time / 呈報日期及時間','____ / ____ / ______   ________'),('Reported by / 呈報人','Name 姓名: __________________  Role 職位: __________________'),('Work activity / 工作活動','____________________________________________________________'),('Employer / contractor / 僱主／承建商','____________________________________________________________'),('Weather / lighting / 天氣／照明','____________________________________________________________'),('Shift / 班次','☐ Day 日班   ☐ Night 夜班   ☐ Other 其他: __________')])
subheading(doc,'Event classification','事件分類')
add_matrix(doc,['Severity / 嚴重程度','Tick / 選擇','Potential severity / 潛在嚴重程度','Tick / 選擇'],[
('Minor / 輕微','☐','Minor / 輕微','☐'),('Medical treatment / 需醫療處理','☐','Serious / 嚴重','☐'),('Lost time / 停工損失','☐','Major / 重大','☐'),('Serious / fatal / 嚴重／致命','☐','Catastrophic / 災難性','☐')],[3800,880,3800,880])
subheading(doc,'Brief factual summary','事件簡述（只列事實）')
add_field_table(doc,[('What happened? / 發生了甚麼？','\n\n\n'),('Immediate outcome / 即時結果','\n\n')])

page_break(doc); heading(doc,'2. PEOPLE INVOLVED','2. 涉事人員')
add_matrix(doc,['Name / 姓名','Role & employer / 職位及僱主','Contact / 聯絡資料','Involvement / 涉事身份'],[
('','','','☐ Injured 傷者  ☐ Witness 證人\n☐ Supervisor 主管  ☐ Other 其他'),('','','','☐ Injured 傷者  ☐ Witness 證人\n☐ Supervisor 主管  ☐ Other 其他'),('','','','☐ Injured 傷者  ☐ Witness 證人\n☐ Supervisor 主管  ☐ Other 其他')],[1800,2500,2000,3060])
subheading(doc,'Injury / illness details (if applicable)','受傷／疾病資料（如適用）')
add_field_table(doc,[('Affected person / 傷病者','__________________________'),('Nature of injury / 傷病性質','____________________________________________________________'),('Body part affected / 受影響身體部位','____________________________________________________________'),('First aid given by / 急救員','__________________________  Time 時間: ________'),('Further treatment / 進一步治療','☐ None 無  ☐ Clinic 診所  ☐ Hospital 醫院  ☐ Ambulance 救護車'),('Work status / 工作狀況','☐ Returned 復工  ☐ Restricted duties 有限度工作  ☐ Lost time 停工'),('Consent/privacy note / 同意／私隱備註','____________________________________________________________')])
subheading(doc,'Witness accounts','證人陳述')
add_field_table(doc,[('Witness 1 / 證人一','Name 姓名: __________________\nStatement 陳述:\n\n'),('Witness 2 / 證人二','Name 姓名: __________________\nStatement 陳述:\n\n')])

page_break(doc); heading(doc,'3. IMMEDIATE RESPONSE','3. 即時應對')
add_matrix(doc,['Time / 時間','Action taken / 採取行動','By whom / 負責人','Status / 狀態'],[('','','','☐ Done 已完成'),('','','','☐ Done 已完成'),('','','','☐ Done 已完成'),('','','','☐ Done 已完成')],[1000,4400,1900,2060])
subheading(doc,'Scene control and notifications','現場管制及通知')
add_field_table(doc,[('Area made safe / 現場已確保安全','☐ Yes 是  ☐ No 否  ☐ Not applicable 不適用   Details 詳情: __________________'),('Work stopped / 已停止工作','☐ Yes 是  ☐ No 否   By whom 由誰: __________________'),('Evidence preserved / 證據已保存','☐ Photos 相片  ☐ CCTV 閉路電視  ☐ Equipment 設備  ☐ Documents 文件'),('Internal notifications / 內部通知','☐ Supervisor 主管  ☐ H&S 安全健康部  ☐ Management 管理層  ☐ HR 人力資源'),('External notifications / 外部通知','☐ Emergency services 緊急服務  ☐ Client 客戶  ☐ Insurer 保險公司  ☐ Authority 主管機關'),('Notification reference / 通知編號','____________________________________________________________')])
note(doc,'Do not disturb the scene unless necessary to rescue a person, remove an immediate danger, or as directed by an authority.\n除非為救人、消除即時危險或依主管機關指示，否則不應擾亂現場。',AMBER)
heading(doc,'4. HAZARD & RISK ASSESSMENT','4. 危害及風險評估')
add_matrix(doc,['Hazard / 危害','Who may be harmed / 可能受影響人士','Existing controls / 現有控制措施','Risk rating / 風險評級'],[('','','','L __ × S __ = __'),('','','','L __ × S __ = __'),('','','','L __ × S __ = __')],[2100,2100,3300,1860])
note(doc,'Suggested scale / 建議評級：Likelihood 可能性 1–5 × Severity 嚴重性 1–5. Apply your organisation’s approved risk matrix where available. / 如機構已有核准風險矩陣，請以該矩陣為準。')

page_break(doc); heading(doc,'5. INVESTIGATION','5. 事故調查')
subheading(doc,'Sequence of events','事件時序')
add_matrix(doc,['Date / time\n日期／時間','Event, action or observation\n事件、行動或觀察','Evidence reference\n證據編號'],[('','',''),('','',''),('','',''),('','',''),('','','')],[1500,5860,2000])
subheading(doc,'Evidence reviewed','已審閱證據')
add_matrix(doc,['Tick / 選擇','Evidence / 證據','Reference / location / 編號／位置'],[
('☐','Photographs / 相片',''),('☐','CCTV / video / 閉路電視／影片',''),('☐','Permits, RAMS, procedures / 許可證、風險評估及工作方法、程序',''),('☐','Training / competency records / 培訓／能力記錄',''),('☐','Inspection / maintenance records / 檢查／維修記錄',''),('☐','Witness statements / 證人陳述',''),('☐','Other / 其他','')],[700,4700,3960])
subheading(doc,'Causal analysis','原因分析')
add_field_table(doc,[('Immediate causes / 直接原因','Unsafe acts or conditions directly preceding the event / 事件前直接出現的不安全行為或情況：\n\n'),('Underlying causes / 潛在原因','Work design, supervision, competency, equipment, environment, communication / 工作設計、監督、能力、設備、環境、溝通：\n\n'),('Root / systemic causes / 根本／系統性原因','Management system, planning, culture, resources, assurance / 管理制度、規劃、文化、資源、保證：\n\n')])

page_break(doc); heading(doc,'6. FINDINGS & CORRECTIVE ACTIONS','6. 調查結果及糾正行動')
subheading(doc,'Key findings','主要調查結果')
add_field_table(doc,[('Finding 1 / 調查結果一','\n'),('Finding 2 / 調查結果二','\n'),('Finding 3 / 調查結果三','\n')])
subheading(doc,'Corrective and preventive action plan','糾正及預防行動計劃')
add_matrix(doc,['No.\n編號','Action / 行動','Hierarchy*\n控制層級*','Owner\n負責人','Due date\n到期日','Status / evidence\n狀態／證據'],[
('1','','','','',''),('2','','','','',''),('3','','','','',''),('4','','','','',''),('5','','','','','')],[550,3200,1450,1250,1100,1810])
note(doc,'* Hierarchy of controls / 控制層級：Eliminate 消除 → Substitute 取代 → Engineering 工程控制 → Administrative 行政控制 → PPE 個人防護裝備。')
subheading(doc,'Residual risk after actions','完成行動後的剩餘風險')
add_field_table(doc,[('Revised rating / 修訂評級','Likelihood 可能性: ____   Severity 嚴重性: ____   Total 總分: ____'),('Acceptable? / 是否可接受？','☐ Yes 是   ☐ No 否   Approved by 批准人: __________________'),('Further monitoring / 進一步監察','____________________________________________________________')])

page_break(doc); heading(doc,'7. NOTIFICATION, REVIEW & CLOSE-OUT','7. 呈報、覆核及結案')
subheading(doc,'Regulatory / contractual reporting','法定／合約呈報')
add_field_table(doc,[('Reportable event? / 是否須呈報？','☐ Yes 是  ☐ No 否  ☐ To be confirmed 待確認'),('Applicable requirement / 適用規定','____________________________________________________________'),('Authority / client notified / 已通知主管機關／客戶','____________________________________________________________'),('Date / time / 日期／時間','____ / ____ / ______   ________'),('Reference / acknowledgement / 編號／確認','____________________________________________________________')])
note(doc,'This template does not determine legal reporting duties. Confirm applicable local law, contract requirements, and reporting deadlines with a competent person.\n本範本不構成法定呈報責任的判定。請由合資格人士確認適用的本地法例、合約要求及呈報期限。',RED)
subheading(doc,'Lessons learned and communication','經驗教訓及傳達')
add_field_table(doc,[('Lessons learned / 經驗教訓','\n\n'),('Communicated to / 已傳達予','☐ Team 團隊  ☐ Contractors 承建商  ☐ Other sites 其他工地  ☐ Client 客戶'),('Method / 方式','☐ Toolbox talk 工具箱講座  ☐ Briefing 簡報  ☐ Alert 通告  ☐ Procedure update 程序更新'),('Date / owner / 日期／負責人','__________________________')])
subheading(doc,'Verification and approval','驗證及批准')
add_matrix(doc,['Role / 角色','Name / 姓名','Signature / 簽署','Date / 日期'],[
('Investigator / 調查員','','',''),('Safety representative / 安全代表','','',''),('Responsible manager / 負責經理','','',''),('Close-out approver / 結案批准人','','','')],[2500,2300,2700,1860])
add_field_table(doc,[('All actions verified? / 所有行動已驗證？','☐ Yes 是  ☐ No 否'),('Closure date / 結案日期','____ / ____ / ______'),('Closure comments / 結案備註','\n\n')])

page_break(doc); heading(doc,'8. ATTACHMENTS & RECORD CONTROL','8. 附件及記錄管制')
add_matrix(doc,['Attach / 附上','Item / 項目','Reference / file name / 編號／檔案名稱'],[
('☐','Photos / site sketch / 相片／現場草圖',''),('☐','Witness statements / 證人陳述',''),('☐','Medical / first-aid record (restricted) / 醫療／急救記錄（受限）',''),('☐','Risk assessment / method statement / 風險評估／工作方法',''),('☐','Permit / inspection / maintenance record / 許可證／檢查／維修記錄',''),('☐','Training / competency record / 培訓／能力記錄',''),('☐','Correspondence / notification receipt / 通訊／通知收據',''),('☐','Other / 其他','')],[900,4600,3860])
subheading(doc,'Document control','文件管制')
add_matrix(doc,['Version / 版本','Date / 日期','Change summary / 修訂摘要','Author / 編製人','Approver / 批准人'],[('','','','',''),('','','','',''),('','','','','')],[1000,1200,3800,1680,1680])
note(doc,'Record retention and privacy / 記錄保存及私隱：Store this report and sensitive personal information in accordance with organisational policy and applicable privacy law. Limit access to authorised persons. / 請依機構政策及適用私隱法例保存本報告及敏感個人資料，並只供獲授權人士查閱。')
subheading(doc,'Optional site sketch','可選現場草圖')
add_field_table(doc,[('Sketch / 草圖','\n\n\n\n\n\n')])

# metadata and core properties
doc.core_properties.title='Bilingual Safety Report Template (English / Traditional Chinese)'
doc.core_properties.subject='Workplace and site safety reporting template'
doc.core_properties.author='Varadise'
doc.core_properties.keywords='safety report, incident, near miss, hazard, English, Traditional Chinese'
doc.save(OUT)
print(OUT)
